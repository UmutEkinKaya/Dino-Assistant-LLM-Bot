"""Cevapları dosyaya döker: PDF, Word, Excel.

Tasarım: dosyanın İÇERİĞİNİ model yazmaz. Model sadece "kullanıcı dosya istedi, şu formatta, şu
kapsamda" der (export_document aracı); içerik, zaten doğrulanmış katmanlı cevaplardan ve kaynak
bilgisinden KODLA üretilir. Böylece dosyada uydurma bir cümle ya da sahte kaynak olamaz.

Bir "tur" sözlüğü: {"question": str, "answer": LayeredAnswer(V2), "sources": {id: SourceInfo|dict}}
"""
from __future__ import annotations

import re
import uuid
from datetime import datetime
from pathlib import Path

from .config import get_settings

FORMATS = {"pdf": "PDF", "docx": "Word", "xlsx": "Excel"}
NAME_RE = re.compile(r"^dino_\d{8}_\d{6}_[0-9a-f]{6}\.(pdf|docx|xlsx)$")   # indirme ucu bunu doğrular


def export_dir() -> Path:
    d = get_settings().export_dir
    d.mkdir(parents=True, exist_ok=True)
    return d


# ------------------------------------------------------------------------------------------
# Ortak içerik: her format aynı satırlardan beslenir (tek doğruluk kaynağı)
# ------------------------------------------------------------------------------------------
LAYERS = [("document", "Belgede yazanlar"), ("inference", "Çıkarım"),
          ("web_facts", "Canlı kaynaklar (belge dışı)")]


def _src_label(sid: str, sources: dict) -> str:
    m = sources.get(sid)
    if m is None:
        return sid
    m = m if isinstance(m, dict) else m.model_dump()
    if m.get("url"):
        return m["url"]
    where = m.get("file") or m.get("doc") or sid
    return f"{where}, s.{m['page']}" if m.get("page") else f"{where} ({m.get('section', '')})"


def rows(turns: list[dict]) -> list[dict]:
    """Excel için düz tablo: her iddia bir satır."""
    out = []
    for n, t in enumerate(turns, 1):
        a, src = t["answer"], t.get("sources") or {}
        out.append({"no": n, "soru": t["question"], "katman": "Özet", "ifade": a.summary, "kaynak": ""})
        if getattr(a, "document_says", ""):
            out.append({"no": n, "soru": t["question"], "katman": "Belgelerde", "ifade": a.document_says, "kaynak": ""})
        for attr, label in LAYERS:
            for c in getattr(a, attr, []) or []:
                out.append({"no": n, "soru": t["question"], "katman": label, "ifade": c.text,
                            "kaynak": "; ".join(_src_label(s, src) for s in c.sources)})
        for x in getattr(a, "debated", []) or []:
            out.append({"no": n, "soru": t["question"], "katman": "Tartışmalı", "ifade": x, "kaynak": ""})
        for x in a.not_in_document:
            out.append({"no": n, "soru": t["question"], "katman": "Belgede yok", "ifade": x, "kaynak": ""})
        for g in a.general_knowledge:
            out.append({"no": n, "soru": t["question"], "katman": "Genel bilgi (belge dışı)",
                        "ifade": g if isinstance(g, str) else g.text, "kaynak": ""})
    return out


def blocks(turns: list[dict]) -> list[tuple[str, str]]:
    """PDF/Word için sıralı bloklar: (tür, metin). tür: h2 / p / quote / h3 / li."""
    out = []
    for n, t in enumerate(turns, 1):
        a, src = t["answer"], t.get("sources") or {}
        out.append(("h2", f"{n}. {t['question']}"))
        out.append(("p", a.summary))
        if getattr(a, "document_says", ""):
            out.append(("quote", f"Belgelerde: {a.document_says}"))
        for attr, label in LAYERS:
            items = getattr(a, attr, []) or []
            if items:
                out.append(("h3", label))
                out += [("li", f"{c.text}  [{'; '.join(_src_label(s, src) for s in c.sources)}]")
                        for c in items]
        for attr, label in [("debated", "Tartışmalı"), ("not_in_document", "Belgede olmayanlar")]:
            items = getattr(a, attr, []) or []
            if items:
                out.append(("h3", label))
                out += [("li", x) for x in items]
        if a.general_knowledge:
            out.append(("h3", "Genel bilgi (belge dışı, doğrulanmamış)"))
            out += [("li", g if isinstance(g, str) else g.text) for g in a.general_knowledge]
    return out


# ------------------------------------------------------------------------------------------
# Formatlar
# ------------------------------------------------------------------------------------------
def _docx(turns, title, path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title, level=1)
    d.add_paragraph(_footer()).runs[0].font.color.rgb = RGBColor(0x78, 0x71, 0x6C)
    for kind, text in blocks(turns):
        if kind == "h2":
            d.add_heading(text, level=2)
        elif kind == "h3":
            d.add_heading(text, level=3)
        elif kind == "li":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "quote":
            d.add_paragraph(text, style="Intense Quote")
        else:
            d.add_paragraph(text)
    d.save(path)


def _xlsx(turns, title, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Cevaplar"
    cols = [("no", "No", 6), ("soru", "Soru", 38), ("katman", "Katman", 22), ("ifade", "İfade", 70),
            ("kaynak", "Kaynak", 42)]
    ws.append([c[1] for c in cols])
    for r in rows(turns):
        ws.append([r[c[0]] for c in cols])
    head = PatternFill("solid", fgColor="2F6F4F")
    for i, (_, _, w) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        ws.column_dimensions[cell.column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Bilgi")
    info.append([title])
    info.append([_footer()])
    info["A1"].font = Font(bold=True, size=14)
    wb.save(path)


def _pdf_font() -> tuple[str, str]:
    """Türkçe karakterleri (ğ ş ı İ) basabilen bir TTF bul. Yoksa Helvetica (karakterler bozulabilir)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pairs = [("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
             ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
             ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
             ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf")]
    for reg, bold in pairs:
        if Path(reg).exists() and Path(bold).exists():
            if "Dino" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Dino", reg))
                pdfmetrics.registerFont(TTFont("Dino-Bold", bold))
            return "Dino", "Dino-Bold"
    return "Helvetica", "Helvetica-Bold"


def _pdf(turns, title, path):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    reg, bold = _pdf_font()
    acc = colors.HexColor("#2F6F4F")
    st = {"h1": ParagraphStyle("h1", fontName=bold, fontSize=18, leading=22, textColor=acc, spaceAfter=4),
          "h2": ParagraphStyle("h2", fontName=bold, fontSize=13, leading=17, spaceBefore=14, spaceAfter=4),
          "h3": ParagraphStyle("h3", fontName=bold, fontSize=9, leading=12, spaceBefore=8,
                               textColor=colors.HexColor("#78716C")),
          "p": ParagraphStyle("p", fontName=reg, fontSize=10.5, leading=15),
          "quote": ParagraphStyle("q", fontName=reg, fontSize=10, leading=14, leftIndent=10, borderPadding=4,
                                  textColor=colors.HexColor("#57534E"), spaceBefore=4),
          "meta": ParagraphStyle("m", fontName=reg, fontSize=8.5, textColor=colors.HexColor("#78716C"))}
    story = [Paragraph(escape(title), st["h1"]), Paragraph(escape(_footer()), st["meta"]), Spacer(1, 6)]
    pending = []

    def flush():
        if pending:
            story.append(ListFlowable([ListItem(Paragraph(escape(x), st["p"]), leftIndent=12) for x in pending],
                                      bulletType="bullet", start="•", leftIndent=12))
            pending.clear()
    for kind, text in blocks(turns):
        if kind == "li":
            pending.append(text)
            continue
        flush()
        story.append(Paragraph(escape(text), st[kind]))
    flush()
    SimpleDocTemplate(str(path), pagesize=A4, title=title, author="Dino Bilgi Asistanı",
                      leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm).build(story)


def _footer() -> str:
    return (f"Dino Bilgi Asistanı · {datetime.now():%d.%m.%Y %H:%M} · İçerik belgelerden ve doğrulanmış "
            "araç sonuçlarından kodla üretildi.")


def render(turns: list[dict], fmt: str, title: str) -> Path:
    if fmt not in FORMATS:
        raise ValueError(f"Bilinmeyen format: {fmt}")
    if not turns:
        raise ValueError("Dışa aktarılacak cevap yok")
    name = f"dino_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}.{fmt}"
    path = export_dir() / name
    {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx}[fmt](turns, title.strip() or "Dino Bilgi Asistanı", path)
    return path


# ------------------------------------------------------------------------------------------
# Veri tablosu (species_table_file aracı): satırlar KODLA bulunmuş değerler, her birinin kaynağı var
# ------------------------------------------------------------------------------------------
TABLE_NOTE = ("Değerler kodla bulundu: önce saha rehberi tablosu, sonra belgelerden çıkarılmış ve "
              "doğrulanmış değerler, yoksa canlı Wikipedia. Dönem adı, yaş aralığından jeolojik zaman "
              "çizelgesiyle hesaplandı. Boş hücre = kaynakta doğrulanabilir değer yok.")


def _new_name(fmt: str) -> Path:
    return export_dir() / f"dino_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}.{fmt}"


def render_table(title: str, rows: list[dict], missing: list[dict], fmt: str) -> Path:
    if fmt not in FORMATS:
        raise ValueError(f"Bilinmeyen format: {fmt}")
    title = title.strip() or "Dinozor tablosu"
    path = _new_name(fmt)
    cols = list(rows[0].keys())
    cell = lambda v: "" if v is None else str(v)
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Tablo"
        ws.append(cols)
        for r in rows:
            ws.append([r[c] if r[c] is not None else None for c in cols])
        widths = {"Cins": 20, "Dönem": 24, "Milyon yıl önce": 16, "Uzunluk (m)": 13, "Ağırlık (t)": 13,
                  "Kaynak türü": 24, "Kaynak": 48}
        for i, c in enumerate(cols, 1):
            h = ws.cell(row=1, column=i)
            h.font, h.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F6F4F")
            ws.column_dimensions[h.column_letter].width = widths.get(c, 18)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink, c.style = c.value, "Hyperlink"
        ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
        if missing:
            m = wb.create_sheet("Bulunamayanlar")
            m.append(["Aranan", "Neden"])
            for x in missing:
                m.append([x["Aranan"], x["Neden"]])
            m.column_dimensions["A"].width, m.column_dimensions["B"].width = 22, 90
            for c in m[1]:
                c.font = Font(bold=True)
        info = wb.create_sheet("Bilgi")
        for line in [title, _footer(), TABLE_NOTE]:
            info.append([line])
        info["A1"].font = Font(bold=True, size=14)
        info.column_dimensions["A"].width = 120
        wb.save(path)
    elif fmt == "docx":
        from docx import Document
        from docx.shared import Pt
        d = Document()
        d.add_heading(title, level=1)
        d.add_paragraph(_footer())
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            t.rows[0].cells[i].text = c
        for r in rows:
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                cells[i].text = cell(r[c])
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
        if missing:
            d.add_heading("Bulunamayanlar", level=2)
            for x in missing:
                d.add_paragraph(f"{x['Aranan']}: {x['Neden']}", style="List Bullet")
        d.add_paragraph(TABLE_NOTE)
        d.save(path)
    else:
        from xml.sax.saxutils import escape
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        reg, bold = _pdf_font()
        sm = ParagraphStyle("c", fontName=reg, fontSize=8, leading=10)
        hd = ParagraphStyle("h", fontName=bold, fontSize=8, leading=10, textColor=colors.white)
        data = [[Paragraph(escape(c), hd) for c in cols]] + \
               [[Paragraph(escape(cell(r[c])), sm) for c in cols] for r in rows]
        w = {"Kaynak": 7.5 * cm, "Dönem": 3.8 * cm, "Kaynak türü": 3.6 * cm}
        tb = Table(data, colWidths=[w.get(c, 2.6 * cm) for c in cols], repeatRows=1)
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F6F4F")),
                                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6D3D1")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F4")])]))
        story = [Paragraph(escape(title), ParagraphStyle("t", fontName=bold, fontSize=16, leading=20,
                                                         textColor=colors.HexColor("#2F6F4F"))),
                 Paragraph(escape(_footer()), ParagraphStyle("m", fontName=reg, fontSize=8)), Spacer(1, 8), tb]
        if missing:
            story += [Spacer(1, 10), Paragraph("Bulunamayanlar: " + escape(
                "; ".join(f"{x['Aranan']} ({x['Neden']})" for x in missing)), sm)]
        story += [Spacer(1, 6), Paragraph(escape(TABLE_NOTE), sm)]
        SimpleDocTemplate(str(path), pagesize=landscape(A4), title=title, leftMargin=1.5 * cm,
                          rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm).build(story)
    return path
