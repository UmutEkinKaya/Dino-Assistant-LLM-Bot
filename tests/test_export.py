"""Dosya dışa aktarma: üç format Türkçe karakterle üretiliyor mu, API akışı doğru mu, indirme ucu güvenli mi."""
import pymupdf as fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from dino import api, export
from dino.schemas import LayeredAnswerV2


def _turn(q="Ceratosaurus Allosaurus'tan büyük müydü?"):
    a = LayeredAnswerV2.model_validate({
        "on_topic": True, "corrections": [], "coverage": "partial", "debated": [],
        "document_says": "Belgeler Ceratosaurus'u yalnızca orta büyüklükte diye tanımlıyor.",
        "document": [{"text": "Ceratosaurus orta büyüklükte bir yırtıcıydı.", "sources": ["ceratosauria:3"]}],
        "inference": [], "web_facts": [{"text": "Ceratosaurus yaklaşık 6–6,7 m uzunluğundaydı.",
                                        "sources": ["web:Ceratosaurus"]}],
        "not_in_document": ["Kesin ağırlık"], "general_knowledge": [],
        "summary": "Hayır; Wikipedia'ya göre Allosaurus daha büyüktü (ğüşıöç İ)."})
    src = {"ceratosauria:3": {"doc": "Ceratosauria", "section": "Description", "file": "Dinozorlar.pdf", "page": 241},
           "web:Ceratosaurus": {"doc": "Ceratosaurus", "section": "Canlı Wikipedia",
                                "url": "https://en.wikipedia.org/wiki/Ceratosaurus"}}
    return {"question": q, "answer": a, "sources": src}


def test_three_formats_with_turkish(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "export_dir", lambda: tmp_path)
    turns = [_turn(), _turn("Peki onu yemiş miydi?")]
    pdf = export.render(turns, "pdf", "Karşılaştırma")
    text = "".join(p.get_text() for p in fitz.open(pdf))
    assert "ğüşıöç" in text and "Dinozorlar.pdf, s.241" in text and "Peki onu yemiş miydi?" in text
    doc = "\n".join(p.text for p in Document(export.render(turns, "docx", "Karşılaştırma")).paragraphs)
    assert "Belgelerde: Belgeler Ceratosaurus" in doc and "https://en.wikipedia.org/wiki/Ceratosaurus" in doc
    ws = load_workbook(export.render(turns, "xlsx", "K"))["Cevaplar"]
    vals = [[c.value for c in r] for r in ws.iter_rows()]
    assert vals[0] == ["No", "Soru", "Katman", "İfade", "Kaynak"]
    assert ["Canlı kaynaklar (belge dışı)", "https://en.wikipedia.org/wiki/Ceratosaurus"] in [[r[2], r[4]] for r in vals]
    assert {r[0] for r in vals[1:]} == {1, 2}


def test_api_export_flow(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "export_dir", lambda: tmp_path)
    monkeypatch.setattr(api, "_sources_meta", lambda ans, links=None: _turn()["sources"])
    api.SESSIONS.clear(); api.ANSWERS.clear()
    monkeypatch.setattr(api.app.state, "retriever", None, raising=False)   # indeks gerekmez: agent sahte
    calls = []

    def fake_ask(q, ret, history=None, allow_export=False, **kw):
        calls.append(allow_export)
        export_turn = "pdf" in q.lower()
        return {"answer": _turn()["answer"], "links": {}, "standalone_question": None, "tickets_opened": [],
                "tools_used": ["export_document"] if export_turn else ["search_docs"], "warnings": [],
                "steps": 1, "trace": [], "latency_ms": 5, "only_export": export_turn, "files": [],
                "export_request": {"format": "pdf", "scope": "conversation", "title": "Sohbet"} if export_turn else None}
    monkeypatch.setattr(api.agent, "ask", fake_ask)
    c = TestClient(api.app)
    sid = "test_session_1"
    c.post("/ask", json={"question": "Ceratosaurus büyük müydü?", "session_id": sid})
    r = c.post("/ask", json={"question": "Bunu PDF yap", "session_id": sid}).json()
    f = r["files"][0]
    assert f["format"] == "pdf" and f["turns"] == 1                   # "PDF yap" turu dosyaya girmez
    dl = c.get(f["url"])
    assert dl.status_code == 200 and dl.content[:4] == b"%PDF"
    # oturumsuz istekte dosya aracı hiç verilmez
    c.post("/ask", json={"question": "Merhaba dinozor"})
    assert calls == [True, True, False]
    # indirme ucu sadece kendi ürettiği isimleri verir
    assert c.get("/exports/..%2F..%2F.env").status_code == 404
    assert c.get("/exports/secret.pdf").status_code == 404


def test_period_from_ages():
    from dino.tools import period_of
    assert period_of(68, 66) == "Geç Kretase"
    assert period_of(150, 140) == "Geç Jura – Erken Kretase"
    assert period_of(float("nan"), 66) is None


def test_species_table_file(tmp_path, monkeypatch):
    from dino import tools, wiki_live
    from dino.tools import SpeciesTableFile, ToolContext, species_table_file
    monkeypatch.setattr(export, "export_dir", lambda: tmp_path)
    monkeypatch.setattr(tools.get_settings(), "live_tools", True)
    live = {"Centrosaurus": {"name": "Centrosaurus", "start_ma": 76.5, "end_ma": 75.5, "length_m_min": 5.5,
                             "length_m_max": 5.5, "mass_t_min": 2.3, "mass_t_max": 2.3, "section": "web:Centrosaurus",
                             "origin": "wikipedia-live", "url": "https://en.wikipedia.org/wiki/Centrosaurus"}}
    monkeypatch.setattr(wiki_live, "lookup_taxon",
                        lambda n: live.get(n, {"error": f"No Wikipedia article named '{n}'."}))
    tools._LIVE_ERRORS.clear()
    a = SpeciesTableFile(names=["Triceratops", "Centrosaurus", "Ceratops", "Nopeasaurus", "triceratops"],
                         format="xlsx", title="Ceratopsidae cinsleri")
    assert "search_docs" in species_table_file(a, ToolContext(allow_files=True))["error"]   # önce belge
    ctx = ToolContext(allow_files=True, searched=True)
    r = species_table_file(a, ctx)
    assert r["rows"] == 2 and set(r["not_found"]) == {"Ceratops", "Nopeasaurus"}   # 'Ceratops' ≠ Triceratops
    wb = load_workbook(tmp_path / r["file"])
    rows = {x[0]: x for x in wb["Tablo"].iter_rows(min_row=2, values_only=True)}
    assert rows["Centrosaurus"][1] == "Geç Kretase" and rows["Centrosaurus"][5] == "Canlı Wikipedia"
    assert rows["Triceratops"][5] == "Rehber tablosu"
    assert {x[0] for x in wb["Bulunamayanlar"].iter_rows(min_row=2, values_only=True)} == {"Ceratops", "Nopeasaurus"}
    assert ctx.files[0]["turns"] == 2
    for fmt in ("docx", "pdf"):                             # diğer formatlar da üretiliyor
        a.format = fmt
        assert species_table_file(a, ToolContext(allow_files=True, searched=True))["status"] == "created"


def _api_with_fake_agent(monkeypatch, tmp_path, calls, files_for=()):
    monkeypatch.setattr(export, "export_dir", lambda: tmp_path)
    monkeypatch.setattr(api, "_sources_meta", lambda ans, links=None: _turn()["sources"])
    monkeypatch.setattr(api.app.state, "retriever", None, raising=False)
    api.SESSIONS.clear(); api.ANSWERS.clear()

    def fake_ask(q, ret, history=None, allow_export=False, on_event=None, **kw):
        calls.append(q)
        if on_event:
            on_event("Belgelerde aranıyor: “Ceratosaurus size”")
        made = [{"name": "dino_20261030_000000_abcdef.xlsx", "format": "xlsx", "turns": 3}] if q in files_for else []
        return {"answer": _turn()["answer"], "links": {}, "standalone_question": None, "tickets_opened": [],
                "tools_used": ["search_docs"], "warnings": [], "steps": 1, "trace": [], "latency_ms": 5,
                "only_export": False, "files": made, "export_request": None,
                "timeline": [{"t": 0.0, "text": "Soru inceleniyor"}, {"t": 1.2, "text": "Belgelerde aranıyor"}]}
    monkeypatch.setattr(api.agent, "ask", fake_ask)
    monkeypatch.setattr(api.agent, "condense",                       # sahte netleştirme (LLM yok)
                        lambda q, h: {"Peki onu yemiş miydi?": "Did Ceratosaurus eat Allosaurus?"}.get(q, q))
    return TestClient(api.app)


def test_answer_cache(tmp_path, monkeypatch):
    calls = []
    c = _api_with_fake_agent(monkeypatch, tmp_path, calls, files_for={"Tablo yap"})
    a = c.post("/ask", json={"question": "Was Ceratosaurus bigger than Allosaurus?"}).json()
    b = c.post("/ask", json={"question": "  was ceratosaurus bigger than allosaurus "}).json()
    assert len(calls) == 1 and not a["cached"] and b["cached"] and b["usage"]["calls"] == 0
    assert b["answer"]["summary"] == a["answer"]["summary"]
    # sohbetin ortasında (geçmişli) aynı soru: netleştirilmiş hali aynı → yine önbellekten
    h = [{"question": "x", "answer": "y"}]
    d = c.post("/ask", json={"question": "Was Ceratosaurus bigger than Allosaurus?", "history": h}).json()
    assert len(calls) == 1 and d["cached"]
    # takip sorusu: anahtar, netleştirilmiş soru ("peki onu..." değil)
    e = c.post("/ask", json={"question": "Peki onu yemiş miydi?", "history": h}).json()
    f = c.post("/ask", json={"question": "Peki onu yemiş miydi?", "history": h}).json()
    assert len(calls) == 2 and not e["cached"] and f["cached"]
    assert f["standalone_question"] == "Did Ceratosaurus eat Allosaurus?"
    # dosya üreten cevap önbelleğe alınmaz
    c.post("/ask", json={"question": "Tablo yap", "session_id": "sess_cache_1"})
    c.post("/ask", json={"question": "Tablo yap", "session_id": "sess_cache_1"})
    assert calls.count("Tablo yap") == 2


def test_stream_sends_steps_then_result(tmp_path, monkeypatch):
    import json as _json
    calls = []
    c = _api_with_fake_agent(monkeypatch, tmp_path, calls)
    with c.stream("POST", "/ask/stream", json={"question": "Was Ceratosaurus bigger?"}) as r:
        assert r.headers["content-type"].startswith("application/x-ndjson")
        events = [_json.loads(line) for line in r.iter_lines() if line.strip()]
    assert [e["type"] for e in events] == ["step", "result"]
    assert events[0]["text"].startswith("Belgelerde aranıyor")
    assert events[1]["data"]["timeline"][1] == {"t": 1.2, "text": "Belgelerde aranıyor"}
    bad = c.post("/ask/stream", json={"question": ""})            # doğrulama hatası akıştan önce: 422
    assert bad.status_code == 422
